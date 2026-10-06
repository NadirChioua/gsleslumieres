from __future__ import annotations

import json
import re
from collections import OrderedDict
from dataclasses import dataclass
from html import unescape
from html.parser import HTMLParser
from pathlib import Path
from typing import Iterable

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "out"
OUTPUT_DIR = ROOT / "deliverables"
OUTPUT_DOCX = OUTPUT_DIR / "revision-textes-site-les-lumieres.docx"

TODAY_FR = "27 août 2026"
SITE_LABEL = "Groupe Scolaire Les Lumières"
SITE_URL = "https://gsleslumieres.ma"

PAGE_ORDER = [
    "/",
    "/qui-sommes-nous/",
    "/mot-du-directeur/",
    "/pourquoi-les-lumieres/",
    "/notre-equipe/",
    "/nos-resultats/",
    "/maternelle-tanger/",
    "/primaire-prive-tanger/",
    "/college-prive-tanger/",
    "/lycee-prive-tanger/",
    "/activites-parascolaires/",
    "/galerie/",
    "/actualites/",
    "/transport-scolaire/",
    "/cantine/",
    "/inscription-ecole-tanger/",
    "/faq/",
    "/contact/",
    "/mentions-legales/",
    "/404/",
]

PAGE_LABELS = {
    "/": "Page d'accueil",
    "/qui-sommes-nous/": "Qui sommes-nous",
    "/mot-du-directeur/": "Mot du Directeur",
    "/pourquoi-les-lumieres/": "Pourquoi Les Lumières",
    "/notre-equipe/": "Notre équipe",
    "/nos-resultats/": "Nos résultats",
    "/maternelle-tanger/": "Maternelle Tanger",
    "/primaire-prive-tanger/": "Primaire privé Tanger",
    "/college-prive-tanger/": "Collège privé Tanger",
    "/lycee-prive-tanger/": "Lycée privé Tanger",
    "/activites-parascolaires/": "Activités parascolaires",
    "/galerie/": "Galerie",
    "/actualites/": "Actualités",
    "/transport-scolaire/": "Transport scolaire",
    "/cantine/": "Cantine",
    "/inscription-ecole-tanger/": "Inscription école Tanger",
    "/faq/": "FAQ",
    "/contact/": "Contact",
    "/mentions-legales/": "Mentions légales",
    "/404/": "Page 404",
}

VOID_TAGS = {
    "area",
    "base",
    "br",
    "col",
    "embed",
    "hr",
    "img",
    "input",
    "link",
    "meta",
    "param",
    "source",
    "track",
    "wbr",
}

SKIP_TAGS = {"script", "style", "svg", "noscript", "template"}
PRIMARY_TEXT_TAGS = {
    "h1",
    "h2",
    "h3",
    "h4",
    "h5",
    "h6",
    "p",
    "li",
    "label",
    "figcaption",
    "dt",
    "dd",
    "option",
}
ACTION_TEXT_TAGS = {"a", "button"}
INLINE_TEXT_TAGS = {"span", "small", "time", "strong", "em"}
TEXT_ATTRS = {"alt", "aria-label", "placeholder", "title"}

TAG_LABELS = {
    "h1": "Titre H1",
    "h2": "Titre de section",
    "h3": "Sous-titre",
    "h4": "Sous-titre",
    "h5": "Sous-titre",
    "h6": "Sous-titre",
    "p": "Paragraphe",
    "li": "Liste",
    "label": "Libellé de champ",
    "figcaption": "Légende média",
    "dt": "Terme",
    "dd": "Détail",
    "option": "Option formulaire",
    "a": "Lien / bouton",
    "button": "Bouton",
    "span": "Libellé court",
    "small": "Note courte",
    "time": "Date",
    "strong": "Mise en avant",
    "em": "Mise en avant",
    "div": "Bloc texte",
}


@dataclass
class Segment:
    kind: str
    text: str


class Node:
    def __init__(self, tag: str = "document", attrs: Iterable[tuple[str, str | None]] | None = None):
        self.tag = tag
        self.attrs = OrderedDict((k, v or "") for k, v in (attrs or []))
        self.children: list[Node] = []
        self.content: list[str | Node] = []

    def text(self) -> str:
        parts: list[str] = []
        for item in self.content:
            if isinstance(item, Node):
                parts.append(item.text())
            else:
                parts.append(item)
        return normalize_text("".join(parts))

    def direct_text(self) -> str:
        return normalize_text("".join(item for item in self.content if isinstance(item, str)))


class TreeBuilder(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.root = Node()
        self.stack = [self.root]

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]):
        node = Node(tag.lower(), attrs)
        self.stack[-1].children.append(node)
        self.stack[-1].content.append(node)
        if tag.lower() not in VOID_TAGS:
            self.stack.append(node)

    def handle_endtag(self, tag: str):
        tag = tag.lower()
        for index in range(len(self.stack) - 1, 0, -1):
            if self.stack[index].tag == tag:
                self.stack = self.stack[:index]
                return

    def handle_data(self, data: str):
        self.stack[-1].content.append(data)


def normalize_text(value: str) -> str:
    value = unescape(value)
    value = value.replace("\u00a0", " ")
    value = value.replace("\u202f", " ")
    value = value.replace("’", "'")
    value = value.replace("–", "-")
    value = value.replace("—", "-")
    value = re.sub(r"\s+", " ", value).strip()
    return value


def meaningful_text(value: str) -> bool:
    if not value or len(value) < 2:
        return False
    if value in {"-", "→", "•", "|", "×", "✓"}:
        return False
    if re.fullmatch(r"[\W_]+", value):
        return False
    return True


def parse_html(path: Path) -> Node:
    parser = TreeBuilder()
    parser.feed(path.read_text(encoding="utf-8"))
    return parser.root


def find_all(node: Node, tag: str) -> list[Node]:
    matches = [node] if node.tag == tag else []
    for child in node.children:
        matches.extend(find_all(child, tag))
    return matches


def find_first(node: Node, tag: str) -> Node | None:
    if node.tag == tag:
        return node
    for child in node.children:
        result = find_first(child, tag)
        if result is not None:
            return result
    return None


def find_direct_child(node: Node, tag: str, *, last: bool = False) -> Node | None:
    children = [child for child in node.children if child.tag == tag]
    if not children:
        return None
    return children[-1] if last else children[0]


def has_descendant(node: Node, tags: set[str]) -> bool:
    for child in node.children:
        if child.tag in tags or has_descendant(child, tags):
            return True
    return False


def element_label(node: Node) -> str:
    return TAG_LABELS.get(node.tag, node.tag)


def extract_segments(node: Node, segments: list[Segment] | None = None) -> list[Segment]:
    if segments is None:
        segments = []
    if node.tag in SKIP_TAGS:
        return segments

    if node.tag in PRIMARY_TEXT_TAGS:
        text = node.text()
        if meaningful_text(text):
            segments.append(Segment(element_label(node), text))
        return segments

    if node.tag in ACTION_TEXT_TAGS:
        if not has_descendant(node, PRIMARY_TEXT_TAGS):
            text = node.text()
            if meaningful_text(text):
                segments.append(Segment(element_label(node), text))
            return segments
        for child in node.children:
            extract_segments(child, segments)
        return segments

    if node.tag in INLINE_TEXT_TAGS and not node.children:
        text = node.text()
        if meaningful_text(text):
            segments.append(Segment(element_label(node), text))
        return segments

    if not node.children:
        text = node.direct_text()
        if meaningful_text(text):
            segments.append(Segment(element_label(node), text))
        return segments

    for child in node.children:
        extract_segments(child, segments)
    return segments


def collect_text_attributes(node: Node, inherited_skip: bool = False) -> list[Segment]:
    skip = inherited_skip or node.tag in SKIP_TAGS
    attrs: list[Segment] = []
    if not skip:
        for attr, value in node.attrs.items():
            if attr in TEXT_ATTRS:
                text = normalize_text(value)
                if meaningful_text(text):
                    attrs.append(Segment(f"{TAG_LABELS.get(node.tag, node.tag)} - {attr}", text))
        for child in node.children:
            attrs.extend(collect_text_attributes(child, skip))
    return attrs


def consolidate(segments: Iterable[Segment]) -> list[tuple[str, str, int]]:
    found: OrderedDict[str, list[object]] = OrderedDict()
    for segment in segments:
        text = normalize_text(segment.text)
        if not meaningful_text(text):
            continue
        if text not in found:
            found[text] = [segment.kind, 0]
        found[text][1] = int(found[text][1]) + 1
    rows: list[tuple[str, str, int]] = []
    for text, (kind, count) in found.items():
        rows.append((str(kind), text, int(count)))
    return rows


def route_from_html(path: Path) -> str:
    relative = path.relative_to(OUT_DIR)
    parts = relative.parts
    if parts == ("index.html",):
        return "/"
    return "/" + "/".join(parts[:-1]) + "/"


def page_sort_key(path: str) -> tuple[int, str]:
    try:
        return PAGE_ORDER.index(path), path
    except ValueError:
        return len(PAGE_ORDER), path


def get_meta(root: Node) -> OrderedDict[str, str]:
    meta: OrderedDict[str, str] = OrderedDict()
    title = find_first(root, "title")
    if title and meaningful_text(title.text()):
        meta["Title HTML"] = title.text()

    for node in find_all(root, "meta"):
        name = node.attrs.get("name") or node.attrs.get("property")
        content = node.attrs.get("content", "")
        if not name or not content:
            continue
        if name in {
            "description",
            "keywords",
            "og:title",
            "og:description",
            "twitter:title",
            "twitter:description",
        }:
            label = {
                "description": "Meta description",
                "keywords": "Meta keywords",
                "og:title": "Open Graph title",
                "og:description": "Open Graph description",
                "twitter:title": "Twitter title",
                "twitter:description": "Twitter description",
            }[name]
            text = normalize_text(content)
            if meaningful_text(text) and text not in meta.values():
                meta[label] = text
    return meta


def page_inventory() -> list[dict[str, object]]:
    pages = []
    for html_path in sorted(OUT_DIR.rglob("index.html")):
        route = route_from_html(html_path)
        root = parse_html(html_path)
        main = find_first(root, "main")
        visible = consolidate(extract_segments(main or root))
        attrs = consolidate(collect_text_attributes(main or root))
        pages.append(
            {
                "route": route,
                "label": PAGE_LABELS.get(route, route.strip("/") or "Accueil"),
                "path": str(html_path),
                "meta": get_meta(root),
                "visible": visible,
                "attrs": attrs,
            }
        )
    return sorted(pages, key=lambda page: page_sort_key(str(page["route"])))


def shared_inventory(home_root: Node) -> OrderedDict[str, list[tuple[str, str, int]]]:
    shared: OrderedDict[str, list[tuple[str, str, int]]] = OrderedDict()
    body = find_first(home_root, "body") or home_root
    header = find_direct_child(body, "header")
    footer = find_direct_child(body, "footer", last=True)
    main = find_first(home_root, "main")

    if header:
        shared["En-tête, barre supérieure et navigation"] = consolidate(
            extract_segments(header) + collect_text_attributes(header)
        )
    if footer:
        shared["Pied de page"] = consolidate(extract_segments(footer) + collect_text_attributes(footer))

    outside_main: list[Segment] = []
    for child in body.children:
        if child is main or child.tag in {"header", "footer"}:
            continue
        outside_main.extend(extract_segments(child))
        outside_main.extend(collect_text_attributes(child))
    if outside_main:
        shared["Éléments hors contenu principal"] = consolidate(outside_main)
    return shared


def source_extras(known_texts: set[str]) -> list[tuple[str, str, int]]:
    manual_segments = [
        Segment("Contact - objet email interne", "Nouveau message de contact - Les Lumières"),
        Segment("Contact - validation", "Requis"),
        Segment("Contact - confirmation", "Message envoyé !"),
        Segment(
            "Contact - confirmation",
            "Merci de nous avoir contactés. Nous vous répondrons très rapidement.",
        ),
        Segment("Inscription - objet email interne", "Nouvelle demande d'inscription - Les Lumières"),
        Segment("Inscription - validation", "Ce champ est requis"),
        Segment("Inscription - validation", "Veuillez choisir un niveau"),
        Segment("Inscription - confirmation", "Merci pour votre demande !"),
        Segment(
            "Inscription - confirmation",
            "Notre équipe vous contactera dans les plus brefs délais via WhatsApp.",
        ),
        Segment("Inscription - confirmation", "Confirmer maintenant sur WhatsApp"),
        Segment(
            "Inscription - message WhatsApp généré",
            "Bonjour, je souhaite inscrire l'étudiant [Nom] en [Niveau] au Groupe Scolaire Les Lumières.",
        ),
        Segment("Menu mobile", "Ouvrir le menu"),
        Segment("Menu mobile", "Menu"),
        Segment("Menu mobile", "Fermer le menu"),
        Segment("Menu mobile - accessibilité", "Navigation mobile"),
        Segment("Header - accessibilité", "Navigation principale"),
        Segment("Header - accessibilité", "Contactez-nous sur WhatsApp"),
        Segment("Bouton WhatsApp flottant", "Contactez-nous sur WhatsApp"),
        Segment("Lecteur vidéo", "Lire la video"),
        Segment("Layout global - accessibilité", "Aller au contenu"),
    ]

    rows = [segment for segment in manual_segments if segment.text not in known_texts]
    return consolidate(rows)


def manifest_inventory() -> list[tuple[str, str, int]]:
    manifest = OUT_DIR / "manifest.webmanifest"
    if not manifest.exists():
        return []
    data = json.loads(manifest.read_text(encoding="utf-8"))
    rows = []
    for key in ("name", "short_name", "description", "start_url", "display", "theme_color", "background_color"):
        value = data.get(key)
        if value is None:
            continue
        text = normalize_text(str(value))
        if meaningful_text(text):
            rows.append((f"Manifest - {key}", text, 1))
    return rows


def set_run_font(run, size: float | None = None, bold: bool | None = None, color: str | None = None):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def set_paragraph_style(paragraph, before: int = 0, after: int = 6, line: float = 1.25):
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = line


def configure_styles(doc: Document):
    styles = doc.styles

    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for name, size, color, before, after in [
        ("Heading 1", 16, "2E74B5", 18, 10),
        ("Heading 2", 13, "2E74B5", 14, 7),
        ("Heading 3", 12, "1F4D78", 10, 5),
    ]:
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.25

    if "Review Small" not in styles:
        small = styles.add_style("Review Small", 1)
    else:
        small = styles["Review Small"]
    small.font.name = "Calibri"
    small._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    small._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    small.font.size = Pt(9)
    small.font.color.rgb = RGBColor.from_string("555555")
    small.paragraph_format.space_after = Pt(4)
    small.paragraph_format.line_spacing = 1.15


def set_section_geometry(doc: Document):
    for section in doc.sections:
        section.start_type = WD_SECTION.NEW_PAGE
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
        section.header_distance = Inches(0.492)
        section.footer_distance = Inches(0.492)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    set_run_font(run, 9, None, "555555")
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_end)


def configure_header_footer(doc: Document):
    section = doc.sections[0]
    header_p = section.header.paragraphs[0]
    header_p.text = ""
    header_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = header_p.add_run("Révision des textes du site web | Groupe Scolaire Les Lumières")
    set_run_font(run, 9, None, "555555")
    set_paragraph_style(header_p, after=3, line=1.0)

    footer_p = section.footer.paragraphs[0]
    footer_p.text = ""
    add_page_number(footer_p)
    set_paragraph_style(footer_p, after=0, line=1.0)


def set_cell_margins(cell, top: int = 80, start: int = 120, bottom: int = 80, end: int = 120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def shade_cell(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.first_child_found_in("w:shd")
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_table_geometry(table, widths: list[float]):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    table.allow_autofit = False
    tbl_pr = table._tbl.tblPr

    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), "9360")
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(int(width * 1440)))
        grid.append(col)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            cell.width = Inches(widths[index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            set_cell_margins(cell)
            tc_w = cell._tc.get_or_add_tcPr().first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                cell._tc.get_or_add_tcPr().append(tc_w)
            tc_w.set(qn("w:w"), str(int(widths[index] * 1440)))
            tc_w.set(qn("w:type"), "dxa")


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def write_cell(cell, text: str, *, bold: bool = False, size: float = 9.3, color: str = "000000"):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    set_paragraph_style(paragraph, after=0, line=1.15)
    run = paragraph.add_run(text)
    set_run_font(run, size, bold, color)


def add_review_table(doc: Document, rows: list[tuple[str, str, int]], empty_message: str = "Aucun texte détecté."):
    table = doc.add_table(rows=1, cols=3)
    set_table_geometry(table, [1.45, 3.55, 1.50])
    headers = ["ÉLÉMENT", "TEXTE ACTUEL", "MODIFICATIONS"]
    for cell, header in zip(table.rows[0].cells, headers):
        shade_cell(cell, "E8EEF5")
        write_cell(cell, header, bold=True, size=9.3, color="0B2545")
    repeat_header(table.rows[0])

    if not rows:
        row = table.add_row()
        write_cell(row.cells[0], "Information")
        write_cell(row.cells[1], empty_message)
        write_cell(row.cells[2], "")
    else:
        for index, (kind, text, count) in enumerate(rows, 1):
            row = table.add_row()
            occurrence = f" ({count} occurrences)" if count > 1 else ""
            write_cell(row.cells[0], f"{kind} {index}{occurrence}")
            write_cell(row.cells[1], text)
            write_cell(row.cells[2], "")

    set_table_geometry(table, [1.45, 3.55, 1.50])
    doc.add_paragraph("", style="Review Small")
    return table


def add_label_detail_table(doc: Document, rows: OrderedDict[str, str] | list[tuple[str, str]]):
    normalized_rows = list(rows.items()) if isinstance(rows, OrderedDict) else rows
    table = doc.add_table(rows=1, cols=2)
    set_table_geometry(table, [1.85, 4.65])
    for cell, header in zip(table.rows[0].cells, ["ÉLÉMENT", "TEXTE ACTUEL"]):
        shade_cell(cell, "E8EEF5")
        write_cell(cell, header, bold=True, size=9.3, color="0B2545")
    repeat_header(table.rows[0])
    for label, text in normalized_rows:
        row = table.add_row()
        write_cell(row.cells[0], label)
        write_cell(row.cells[1], text)
    set_table_geometry(table, [1.85, 4.65])
    doc.add_paragraph("", style="Review Small")
    return table


def add_title(doc: Document, title: str, subtitle: str):
    p = doc.add_paragraph()
    set_paragraph_style(p, before=0, after=3, line=1.0)
    run = p.add_run(title)
    set_run_font(run, 24, True, "0B2545")

    p = doc.add_paragraph()
    set_paragraph_style(p, before=0, after=12, line=1.15)
    run = p.add_run(subtitle)
    set_run_font(run, 12, None, "555555")


def add_note(doc: Document, text: str):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [6.5])
    shade_cell(table.cell(0, 0), "F4F6F9")
    write_cell(table.cell(0, 0), text, size=9.8, color="0B2545")
    doc.add_paragraph("", style="Review Small")


def make_document():
    pages = page_inventory()
    home_root = parse_html(OUT_DIR / "index.html")
    shared = shared_inventory(home_root)

    known_texts = set()
    for group in shared.values():
        known_texts.update(text for _, text, _ in group)
    for page in pages:
        known_texts.update(text for _, text, _ in page["visible"])  # type: ignore[index]
        known_texts.update(text for _, text, _ in page["attrs"])  # type: ignore[index]
        known_texts.update(page["meta"].values())  # type: ignore[index]

    extras = source_extras(known_texts)
    manifest_rows = manifest_inventory()

    doc = Document()
    set_section_geometry(doc)
    configure_styles(doc)
    configure_header_footer(doc)

    add_title(
        doc,
        "Révision complète des textes du site web",
        f"{SITE_LABEL} - Document de travail généré le {TODAY_FR}",
    )
    add_note(
        doc,
        "Objectif : permettre à l'école de relire et corriger tous les textes du site. "
        "La colonne « TEXTE ACTUEL » reprend les textes détectés dans la version actuelle du site, "
        "et la colonne « MODIFICATIONS » est laissée vide pour les corrections, suppressions ou validations."
    )

    add_label_detail_table(
        doc,
        OrderedDict(
            [
                ("Site", SITE_URL),
                ("Version auditée", "Build local Next.js généré avant extraction"),
                ("Nombre de pages HTML", str(len(pages))),
                ("Méthode", "Extraction du HTML rendu + textes SEO + textes de formulaires et messages conditionnels"),
                ("Style du document", "compact_reference_guide"),
            ]
        ),
    )

    doc.add_heading("Sommaire", level=1)
    summary_rows = OrderedDict()
    summary_rows["I. Éléments globaux"] = "En-tête, navigation, pied de page, boutons, formulaires et messages conditionnels"
    for index, page in enumerate(pages, 1):
        route = str(page["route"])
        summary_rows[f"{index}. {page['label']}"] = route
    summary_rows["Annexe"] = "Manifest application et textes techniques restants"
    add_label_detail_table(doc, summary_rows)

    doc.add_heading("I. Éléments globaux", level=1)
    for group_title, rows in shared.items():
        doc.add_heading(group_title, level=2)
        add_review_table(doc, rows)

    doc.add_heading("Boutons, formulaires et messages conditionnels", level=2)
    add_review_table(
        doc,
        extras,
        "Aucun texte conditionnel supplémentaire détecté dans les composants ciblés.",
    )

    for page in pages:
        route = str(page["route"])
        label = str(page["label"])
        doc.add_page_break()
        doc.add_heading(f"{label} - {route}", level=1)

        meta = page["meta"]  # type: ignore[assignment]
        if meta:
            doc.add_heading("Métadonnées SEO", level=2)
            add_label_detail_table(doc, meta)  # type: ignore[arg-type]

        doc.add_heading("Textes visibles dans la page", level=2)
        add_review_table(doc, page["visible"])  # type: ignore[arg-type]

        attrs = page["attrs"]  # type: ignore[assignment]
        if attrs:
            doc.add_heading("Textes d'images, champs et accessibilité", level=2)
            add_review_table(doc, attrs)  # type: ignore[arg-type]

    if manifest_rows:
        doc.add_page_break()
        doc.add_heading("Annexe - Manifest application", level=1)
        add_review_table(doc, manifest_rows)

    OUTPUT_DIR.mkdir(exist_ok=True)
    doc.save(OUTPUT_DOCX)
    return OUTPUT_DOCX, len(pages), sum(len(page["visible"]) for page in pages), sum(len(page["attrs"]) for page in pages)


if __name__ == "__main__":
    output, page_count, visible_count, attr_count = make_document()
    print(f"Created: {output}")
    print(f"Pages inventoried: {page_count}")
    print(f"Visible text rows: {visible_count}")
    print(f"Attribute/accessibility rows: {attr_count}")
